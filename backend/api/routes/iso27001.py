from fastapi import APIRouter, BackgroundTasks, UploadFile, File, HTTPException
from fastapi.responses import FileResponse, JSONResponse
from pydantic import BaseModel
from typing import List, Dict, Any, Optional
from services.chat_service import ChatService
import uuid
import json
import os
import shutil
import logging
from pathlib import Path
from datetime import datetime, timezone

logger = logging.getLogger(__name__)
router = APIRouter()

ASSESSMENTS_DIR = os.getenv("DATA_PATH", "./data") + "/assessments"
EVIDENCE_DIR = os.getenv("DATA_PATH", "./data") + "/evidence"
EXPORTS_DIR = os.getenv("DATA_PATH", "./data") + "/exports"
os.makedirs(ASSESSMENTS_DIR, exist_ok=True)
os.makedirs(EVIDENCE_DIR, exist_ok=True)
os.makedirs(EXPORTS_DIR, exist_ok=True)

ALLOWED_EVIDENCE_EXT = {".pdf", ".png", ".jpg", ".jpeg", ".doc", ".docx", ".xlsx", ".csv", ".txt", ".log", ".conf", ".xml", ".json"}
TEXT_EXTENSIONS = {".txt", ".log", ".conf", ".xml", ".json", ".csv"}
MAX_EVIDENCE_SIZE = 10 * 1024 * 1024  # 10MB per file
MAX_EVIDENCE_CONTENT_CHARS = 3000  # Max chars per file for AI prompt


def parse_evidence_file_content(filepath: str) -> str:
    """Extract text content from evidence file for AI analysis."""
    ext = os.path.splitext(filepath)[1].lower()
    try:
        if ext in TEXT_EXTENSIONS:
            with open(filepath, "r", encoding="utf-8", errors="replace") as f:
                content = f.read(MAX_EVIDENCE_CONTENT_CHARS)
                if len(content) >= MAX_EVIDENCE_CONTENT_CHARS:
                    content += "\n... [truncated]"
                return content
        elif ext == ".pdf":
            try:
                from pypdf import PdfReader
                reader = PdfReader(filepath)
                text = ""
                for page in reader.pages[:5]:  # Max 5 pages
                    text += page.extract_text() or ""
                    if len(text) >= MAX_EVIDENCE_CONTENT_CHARS:
                        break
                return text[:MAX_EVIDENCE_CONTENT_CHARS] + ("\n... [truncated]" if len(text) > MAX_EVIDENCE_CONTENT_CHARS else "")
            except ImportError:
                return "[PDF file — pypdf not installed, content not extracted]"
            except Exception as e:
                return f"[PDF parse error: {str(e)[:100]}]"
        elif ext in {".doc", ".docx"}:
            try:
                from docx import Document
                doc = Document(filepath)
                text = "\n".join([p.text for p in doc.paragraphs[:50]])
                return text[:MAX_EVIDENCE_CONTENT_CHARS]
            except ImportError:
                return "[DOCX file — python-docx not installed]"
            except Exception:
                return "[DOCX file — could not parse]"
        elif ext == ".xlsx":
            try:
                import openpyxl
                wb = openpyxl.load_workbook(filepath, read_only=True)
                text = ""
                for ws in wb.worksheets[:3]:
                    for row in ws.iter_rows(max_row=30, values_only=True):
                        text += " | ".join([str(c) if c else "" for c in row]) + "\n"
                        if len(text) >= MAX_EVIDENCE_CONTENT_CHARS:
                            break
                return text[:MAX_EVIDENCE_CONTENT_CHARS]
            except ImportError:
                return "[XLSX file — openpyxl not installed]"
            except Exception:
                return "[XLSX file — could not parse]"
        elif ext in {".png", ".jpg", ".jpeg"}:
            size_kb = os.path.getsize(filepath) // 1024
            return f"[Image file ({ext}) — {size_kb}KB — visual evidence attached]"
        else:
            return f"[Binary file ({ext}) — content not extractable]"
    except Exception as e:
        logger.warning(f"Evidence parse error {filepath}: {e}")
        return f"[Error reading file: {str(e)[:100]}]"


def build_evidence_context_for_ai(evidence_map: Dict[str, List[str]]) -> str:
    """Build structured evidence text for AI prompt from evidence_map.
    evidence_map: { controlId: [filename1, filename2, ...] }
    """
    if not evidence_map:
        return ""

    sections = []
    sections.append("\n\n══════ BẰNG CHỨNG TRIỂN KHAI (EVIDENCE) ══════")
    for ctrl_id, filenames in evidence_map.items():
        if not filenames:
            continue
        ctrl_dir = os.path.join(EVIDENCE_DIR, ctrl_id.replace(".", "_"))
        if not os.path.exists(ctrl_dir):
            sections.append(f"\n📎 {ctrl_id}: [Thư mục evidence không tồn tại]")
            continue

        sections.append(f"\n📎 {ctrl_id} ({len(filenames)} file minh chứng):")
        for fname in filenames[:5]:  # Max 5 files per control
            fpath = os.path.join(ctrl_dir, fname)
            if os.path.exists(fpath):
                content = parse_evidence_file_content(fpath)
                sections.append(f"  ├─ {fname}:")
                # Indent content
                for line in content.split("\n")[:20]:
                    sections.append(f"  │  {line}")
            else:
                sections.append(f"  ├─ {fname}: [file not found on server]")

    sections.append("══════ KẾT THÚC BẰNG CHỨNG ══════\n")
    sections.append("Lưu ý cho Auditor: Các controls có bằng chứng đính kèm cho thấy ")
    sections.append("tổ chức đã có tài liệu/minh chứng triển khai thực tế. ")
    sections.append("Hãy đánh giá CHẤT LƯỢNG bằng chứng và mức độ đầy đủ.")
    return "\n".join(sections)


class SystemInfo(BaseModel):
    assessment_standard: str = "iso27001"
    org_name: str = ""
    org_size: str = ""
    industry: str = ""
    servers: int = 0
    firewalls: str = ""
    vpn: bool = False
    cloud_provider: str = ""
    antivirus: str = ""
    backup_solution: str = ""
    siem: str = ""
    network_diagram: str = ""
    implemented_controls: List[str] = []
    incidents_12m: int = 0
    employees: int = 0
    it_staff: int = 0
    iso_status: str = ""
    notes: str = ""
    model_mode: str = "hybrid"  # "hybrid" | "local" | "cloud"
    evidence_map: Dict[str, List[str]] = {}  # { controlId: [filename1, ...] }


def save_assessment(assessment_id: str, data: dict):
    filepath = os.path.join(ASSESSMENTS_DIR, f"{assessment_id}.json")
    with open(filepath, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)


def load_assessment(assessment_id: str) -> Optional[dict]:
    filepath = os.path.join(ASSESSMENTS_DIR, f"{assessment_id}.json")
    if os.path.exists(filepath):
        with open(filepath, "r", encoding="utf-8") as f:
            return json.load(f)
    return None


def list_assessments() -> List[dict]:
    results = []
    for filename in os.listdir(ASSESSMENTS_DIR):
        if filename.endswith(".json"):
            filepath = os.path.join(ASSESSMENTS_DIR, filename)
            try:
                with open(filepath, "r", encoding="utf-8") as f:
                    data = json.load(f)
                    results.append({
                        "id": data.get("id"),
                        "status": data.get("status"),
                        "standard": data.get("system_info", {}).get("assessment_standard", "iso27001"),
                        "org_name": data.get("system_info", {}).get("organization", {}).get("name", "Unknown"),
                        "created_at": data.get("created_at"),
                        "updated_at": data.get("updated_at"),
                        "compliance_percent": data.get("compliance_percent")
                    })
            except Exception:
                pass
    return sorted(results, key=lambda x: x.get("created_at", ""), reverse=True)


def process_assessment_bg(assessment_id: str, system_data: dict, model_mode: str = "hybrid", evidence_context: str = ""):
    data = load_assessment(assessment_id)
    if not data:
        return

    data["status"] = "processing"
    data["updated_at"] = datetime.now(timezone.utc).isoformat()
    save_assessment(assessment_id, data)

    try:
        # Append evidence context to notes for AI processing
        if evidence_context:
            system_data["notes"] = (system_data.get("notes", "") or "") + evidence_context
            logger.info(f"[Assessment {assessment_id[:8]}] Evidence context appended ({len(evidence_context)} chars)")

        result = ChatService.assess_system(system_data, model_mode=model_mode)

        data["status"] = "completed"
        data["result"] = result
        # Store json_data at top level for quick access
        if result.get("json_data"):
            data["json_data"] = result["json_data"]
        data["updated_at"] = datetime.now(timezone.utc).isoformat()
        save_assessment(assessment_id, data)

    except Exception as e:
        data["status"] = "failed"
        data["error"] = str(e)
        data["updated_at"] = datetime.now(timezone.utc).isoformat()
        save_assessment(assessment_id, data)


@router.post("/iso27001/assess")
async def assess(data: SystemInfo, background_tasks: BackgroundTasks):
    assessment_id = str(uuid.uuid4())
    system_data = {
        "assessment_standard": data.assessment_standard,
        "organization": {
            "name": data.org_name,
            "size": data.org_size,
            "industry": data.industry,
            "employees": data.employees,
            "it_staff": data.it_staff
        },
        "infrastructure": {
            "servers": data.servers,
            "firewalls": data.firewalls,
            "vpn": "Có" if data.vpn else "Không",
            "cloud": data.cloud_provider or "Không sử dụng",
            "antivirus": data.antivirus or "Không có",
            "backup": data.backup_solution or "Không có",
            "siem": data.siem or "Không có",
            "network_diagram": data.network_diagram or "Không cung cấp"
        },
        "compliance": {
            "iso_status": data.iso_status or "Chưa triển khai",
            "implemented_controls": data.implemented_controls or [],
            "incidents_12m": data.incidents_12m
        },
        "notes": data.notes,
        "model_mode": data.model_mode
    }

    # Build evidence context from parsed file contents
    evidence_context = ""
    if data.evidence_map:
        evidence_context = build_evidence_context_for_ai(data.evidence_map)
        logger.info(f"[Assessment] Evidence map: {len(data.evidence_map)} controls with evidence")

    # Tính compliance_percent sơ bộ ngay lúc tạo (weighted: critical=4, high=3, medium=2, low=1)
    impl_controls = data.implemented_controls or []
    total_controls = 93 if data.assessment_standard == "iso27001" else 34
    compliance_pct = 0
    try:
        from services.standard_service import load_standard, WEIGHT_SCORE
        custom_std = load_standard(data.assessment_standard)
        if custom_std:
            all_ctrls = []
            for cat in custom_std.get("controls", []):
                all_ctrls.extend(cat.get("controls", []))
            total_controls = len(all_ctrls)
            weight_map = {c["id"]: WEIGHT_SCORE.get(c.get("weight", "medium"), 1) for c in all_ctrls}
            max_score = sum(weight_map.values())
            achieved = sum(weight_map.get(cid, 0) for cid in impl_controls)
            compliance_pct = round((achieved / max_score) * 100, 1) if max_score > 0 else 0
        else:
            compliance_pct = round((len(impl_controls) / total_controls) * 100, 1) if total_controls > 0 else 0
    except Exception:
        compliance_pct = round((len(impl_controls) / total_controls) * 100, 1) if total_controls > 0 else 0

    assessment_record = {
        "id": assessment_id,
        "status": "pending",
        "system_info": system_data,
        "compliance_percent": compliance_pct,
        "model_mode": data.model_mode,
        "standard": data.assessment_standard,
        "evidence_attached": len(data.evidence_map) > 0,
        "created_at": datetime.now(timezone.utc).isoformat(),
        "updated_at": datetime.now(timezone.utc).isoformat()
    }

    save_assessment(assessment_id, assessment_record)
    background_tasks.add_task(process_assessment_bg, assessment_id, system_data, data.model_mode, evidence_context)

    return {"status": "accepted", "id": assessment_id, "message": "Assessment task started in background"}


@router.get("/iso27001/assessments")
async def get_all_assessments():
    return list_assessments()


@router.get("/iso27001/assessments/{assessment_id}")
async def get_assessment(assessment_id: str):
    data = load_assessment(assessment_id)
    if not data:
        return {"error": "Assessment not found", "status": "not_found"}
    return data


@router.delete("/iso27001/assessments/{assessment_id}")
async def delete_assessment(assessment_id: str):
    filepath = os.path.join(ASSESSMENTS_DIR, f"{assessment_id}.json")
    if os.path.exists(filepath):
        try:
            os.remove(filepath)
            return {"status": "success", "message": "Assessment deleted successfully"}
        except Exception as e:
            return {"status": "error", "message": f"Failed to delete: {str(e)}"}
    return {"status": "not_found", "message": "Assessment not found"}


@router.post("/iso27001/reindex")
async def reindex():
    vs = ChatService.get_vector_store()
    result = vs.index_documents()
    return result


@router.get("/iso27001/chromadb/stats")
async def chromadb_stats():
    try:
        vs = ChatService.get_vector_store()
        count = vs.collection.count()
        metadata = vs.collection.metadata
        peek = vs.collection.peek(limit=3)
        sources = set()
        if peek and peek.get('metadatas'):
            for m in peek['metadatas']:
                if m and m.get('source'):
                    sources.add(m['source'])

        docs_dir = os.getenv("ISO_DOCS_PATH", "/data/iso_documents")
        files = []
        from pathlib import Path
        docs_path = Path(docs_dir)
        if docs_path.exists():
            for f in docs_path.glob("*.md"):
                files.append({
                    "name": f.name,
                    "size_bytes": f.stat().st_size
                })

        return {
            "status": "ok",
            "total_chunks": count,
            "total_files": len(files),
            "files": files,
            "collection_name": "iso_documents",
            "metric": metadata.get("hnsw:space", "unknown"),
            "sample_sources": list(sources)
        }
    except Exception as e:
        return {"status": "error", "message": str(e)}


@router.post("/iso27001/chromadb/search")
async def chromadb_search(query: dict):
    try:
        vs = ChatService.get_vector_store()
        q = query.get("query", "")
        top_k = query.get("top_k", 3)
        if not q:
            return {"status": "error", "message": "Missing query parameter"}
        results = vs.search(q, top_k=top_k)
        return {"status": "ok", "query": q, "results": results}
    except Exception as e:
        return {"status": "error", "message": str(e)}


# ═══════════════════════════════════════════════════════════
# EVIDENCE UPLOAD — per control
# ═══════════════════════════════════════════════════════════

@router.post("/iso27001/evidence/{control_id}")
async def upload_evidence(control_id: str, file: UploadFile = File(...)):
    """Upload evidence file for a specific control."""
    if not file.filename:
        raise HTTPException(status_code=400, detail="No filename")

    ext = "." + file.filename.rsplit(".", 1)[-1].lower() if "." in file.filename else ""
    if ext not in ALLOWED_EVIDENCE_EXT:
        raise HTTPException(status_code=400, detail=f"File type '{ext}' not allowed. Allowed: {', '.join(ALLOWED_EVIDENCE_EXT)}")

    content = await file.read()
    if len(content) > MAX_EVIDENCE_SIZE:
        raise HTTPException(status_code=413, detail=f"File too large. Max: {MAX_EVIDENCE_SIZE // (1024*1024)}MB")

    # Create control-specific directory
    ctrl_dir = os.path.join(EVIDENCE_DIR, control_id.replace(".", "_"))
    os.makedirs(ctrl_dir, exist_ok=True)

    # Save with timestamp prefix to avoid overwrite
    ts = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
    safe_name = f"{ts}_{file.filename}"
    filepath = os.path.join(ctrl_dir, safe_name)

    with open(filepath, "wb") as f:
        f.write(content)

    return {
        "status": "success",
        "control_id": control_id,
        "filename": safe_name,
        "size_bytes": len(content),
        "path": f"/api/iso27001/evidence/{control_id}/{safe_name}",
    }


@router.get("/iso27001/evidence/{control_id}")
async def list_evidence(control_id: str):
    """List all evidence files for a control."""
    ctrl_dir = os.path.join(EVIDENCE_DIR, control_id.replace(".", "_"))
    if not os.path.exists(ctrl_dir):
        return {"control_id": control_id, "files": []}

    files = []
    for filename in sorted(os.listdir(ctrl_dir)):
        filepath = os.path.join(ctrl_dir, filename)
        stat = os.stat(filepath)
        files.append({
            "filename": filename,
            "size_bytes": stat.st_size,
            "uploaded_at": datetime.fromtimestamp(stat.st_mtime, tz=timezone.utc).isoformat(),
            "download_url": f"/api/iso27001/evidence/{control_id}/{filename}",
        })

    return {"control_id": control_id, "files": files}


@router.get("/iso27001/evidence/{control_id}/{filename}")
async def download_evidence(control_id: str, filename: str):
    """Download a specific evidence file."""
    ctrl_dir = os.path.join(EVIDENCE_DIR, control_id.replace(".", "_"))
    filepath = os.path.join(ctrl_dir, filename)

    if not os.path.exists(filepath):
        raise HTTPException(status_code=404, detail="File not found")

    return FileResponse(filepath, filename=filename)


@router.delete("/iso27001/evidence/{control_id}/{filename}")
async def delete_evidence(control_id: str, filename: str):
    """Delete a specific evidence file."""
    ctrl_dir = os.path.join(EVIDENCE_DIR, control_id.replace(".", "_"))
    filepath = os.path.join(ctrl_dir, filename)

    if not os.path.exists(filepath):
        raise HTTPException(status_code=404, detail="File not found")

    os.remove(filepath)
    return {"status": "success", "message": f"Deleted {filename}"}


@router.get("/iso27001/evidence-summary")
async def get_all_evidence_summary():
    """Get summary of all uploaded evidence across all controls."""
    summary = {}
    if not os.path.exists(EVIDENCE_DIR):
        return {"controls": {}, "total_files": 0}

    total = 0
    for ctrl_folder in os.listdir(EVIDENCE_DIR):
        ctrl_path = os.path.join(EVIDENCE_DIR, ctrl_folder)
        if os.path.isdir(ctrl_path):
            ctrl_id = ctrl_folder.replace("_", ".")
            files = [f for f in os.listdir(ctrl_path) if os.path.isfile(os.path.join(ctrl_path, f))]
            if files:
                summary[ctrl_id] = len(files)
                total += len(files)

    return {"controls": summary, "total_files": total}


@router.get("/iso27001/evidence/{control_id}/{filename}/preview")
async def preview_evidence(control_id: str, filename: str):
    """Parse and return text content of an evidence file for preview."""
    ctrl_dir = os.path.join(EVIDENCE_DIR, control_id.replace(".", "_"))
    filepath = os.path.join(ctrl_dir, filename)

    if not os.path.exists(filepath):
        raise HTTPException(status_code=404, detail="File not found")

    ext = os.path.splitext(filename)[1].lower()
    content = parse_evidence_file_content(filepath)
    is_image = ext in {".png", ".jpg", ".jpeg"}

    return {
        "control_id": control_id,
        "filename": filename,
        "content": content,
        "content_type": "image" if is_image else "text",
        "file_ext": ext,
        "size_bytes": os.path.getsize(filepath),
        "download_url": f"/api/iso27001/evidence/{control_id}/{filename}"
    }


# ═══════════════════════════════════════════════════════════
# PDF EXPORT — server-side
# ═══════════════════════════════════════════════════════════

@router.post("/iso27001/assessments/{assessment_id}/export-pdf")
async def export_pdf(assessment_id: str):
    """Generate PDF from assessment report using weasyprint.
    Falls back to HTML file if weasyprint is not available."""
    data = load_assessment(assessment_id)
    if not data:
        raise HTTPException(status_code=404, detail="Assessment not found")

    if data.get("status") != "completed":
        raise HTTPException(status_code=400, detail="Assessment not completed yet")

    report = data.get("result", {}).get("report", "")
    if not report:
        raise HTTPException(status_code=400, detail="No report content")

    sys_info = data.get("system_info", {})
    org_name = sys_info.get("organization", {}).get("name", "Tổ chức")
    std = data.get("standard", "iso27001")
    std_name = "ISO 27001:2022" if std == "iso27001" else "TCVN 11930:2017" if std == "tcvn11930" else std
    pct = data.get("compliance_percent", 0)
    created = data.get("created_at", "")

    pct_color = '#10b981' if pct >= 80 else '#3b82f6' if pct >= 50 else '#f59e0b' if pct >= 25 else '#ef4444'

    # Convert markdown to HTML (basic conversion)
    import re as _re
    html_body = report
    html_body = _re.sub(r'^#{1}\s+(.+)$', r'<h1>\1</h1>', html_body, flags=_re.MULTILINE)
    html_body = _re.sub(r'^#{2}\s+(.+)$', r'<h2>\1</h2>', html_body, flags=_re.MULTILINE)
    html_body = _re.sub(r'^#{3}\s+(.+)$', r'<h3>\1</h3>', html_body, flags=_re.MULTILINE)
    html_body = _re.sub(r'^#{4}\s+(.+)$', r'<h4>\1</h4>', html_body, flags=_re.MULTILINE)
    html_body = _re.sub(r'\*\*(.+?)\*\*', r'<strong>\1</strong>', html_body)
    html_body = _re.sub(r'\*(.+?)\*', r'<em>\1</em>', html_body)
    html_body = _re.sub(r'^[-•]\s+(.+)$', r'<li>\1</li>', html_body, flags=_re.MULTILINE)
    html_body = _re.sub(r'(<li>.*?</li>\n?)+', lambda m: f'<ul>{m.group()}</ul>', html_body, flags=_re.DOTALL)
    html_body = _re.sub(r'^---+$', '<hr>', html_body, flags=_re.MULTILINE)

    # Handle markdown tables
    def convert_md_table(match):
        lines = match.group().strip().split('\n')
        if len(lines) < 2:
            return match.group()
        headers = [h.strip() for h in lines[0].strip('|').split('|')]
        rows_html = '<tr>' + ''.join(f'<th>{h}</th>' for h in headers) + '</tr>\n'
        for line in lines[2:]:
            cells = [c.strip() for c in line.strip('|').split('|')]
            rows_html += '<tr>' + ''.join(f'<td>{c}</td>' for c in cells) + '</tr>\n'
        return f'<table>{rows_html}</table>'

    html_body = _re.sub(r'(\|.+\|(?:\n\|[-:| ]+\|)(?:\n\|.+\|)+)', convert_md_table, html_body)
    html_body = html_body.replace('\n\n', '</p><p>')

    html_content = f"""<!DOCTYPE html>
<html lang="vi">
<head>
<meta charset="UTF-8">
<title>Báo cáo Đánh giá ATTT - {org_name}</title>
<style>
  @page {{ margin: 2cm; size: A4; }}
  body {{ font-family: 'Segoe UI', Arial, Helvetica, sans-serif; max-width: 780px; margin: 0 auto; color: #1e293b; line-height: 1.7; font-size: 13px; }}
  h1 {{ font-size: 20px; font-weight: 800; border-bottom: 2px solid #3b82f6; padding-bottom: 8px; margin-top: 24px; }}
  h2 {{ font-size: 16px; font-weight: 700; color: #3b82f6; margin-top: 20px; border-left: 3px solid #3b82f6; padding-left: 8px; }}
  h3 {{ font-size: 14px; font-weight: 600; color: #475569; margin-top: 14px; }}
  h4 {{ font-size: 13px; font-weight: 600; color: #64748b; margin-top: 10px; }}
  .hero {{ display: flex; align-items: center; gap: 20px; background: #f8fafc; border: 1px solid #e2e8f0; border-radius: 12px; padding: 16px 20px; margin: 16px 0; }}
  .pct {{ font-size: 36px; font-weight: 900; color: {pct_color}; min-width: 90px; text-align: center; }}
  .meta {{ flex: 1; }}
  .meta strong {{ display: block; font-size: 15px; }}
  .meta span {{ color: #64748b; font-size: 12px; }}
  table {{ width: 100%; border-collapse: collapse; margin: 10px 0; font-size: 12px; }}
  th, td {{ padding: 7px 10px; text-align: left; border: 1px solid #e2e8f0; }}
  th {{ background: #f1f5f9; font-weight: 600; color: #475569; }}
  ul, ol {{ padding-left: 20px; }}
  li {{ margin-bottom: 3px; }}
  strong {{ color: #1e293b; }}
  blockquote {{ border-left: 3px solid #3b82f6; margin: 8px 0; padding: 6px 12px; background: #eff6ff; border-radius: 0 8px 8px 0; color: #475569; font-size: 12px; }}
  hr {{ border: none; border-top: 1px solid #e2e8f0; margin: 14px 0; }}
  code {{ background: #f1f5f9; border: 1px solid #e2e8f0; border-radius: 3px; padding: 1px 5px; font-size: 11px; }}
  .footer {{ margin-top: 30px; padding-top: 12px; border-top: 1px solid #e2e8f0; font-size: 10px; color: #94a3b8; text-align: center; }}
</style>
</head>
<body>
<div class="hero">
  <div class="pct">{pct}%</div>
  <div class="meta">
    <strong>{org_name}</strong>
    <span>{std_name} · {created[:10] if created else 'N/A'} · ID: {assessment_id[:8]}</span>
  </div>
</div>
{html_body}
<div class="footer">
  Báo cáo tự động bởi CyberAI Assessment Platform · {created[:10] if created else ''} · ID: {assessment_id}
</div>
</body>
</html>"""

    pdf_filename = f"report_{assessment_id[:8]}_{org_name.replace(' ', '_')[:30]}.pdf"
    html_filename = f"report_{assessment_id[:8]}.html"
    pdf_path = os.path.join(EXPORTS_DIR, pdf_filename)
    html_path = os.path.join(EXPORTS_DIR, html_filename)

    # Save HTML regardless
    with open(html_path, "w", encoding="utf-8") as f:
        f.write(html_content)

    # Try weasyprint for PDF
    try:
        from weasyprint import HTML
        HTML(string=html_content).write_pdf(pdf_path)
        return FileResponse(
            pdf_path,
            media_type="application/pdf",
            filename=pdf_filename,
        )
    except ImportError:
        # weasyprint not installed — return HTML file
        return FileResponse(
            html_path,
            media_type="text/html",
            filename=html_filename,
            headers={"X-PDF-Fallback": "true", "X-Message": "weasyprint not installed, returning HTML"},
        )
    except Exception as e:
        # weasyprint error — return HTML as fallback
        return FileResponse(
            html_path,
            media_type="text/html",
            filename=html_filename,
            headers={"X-PDF-Fallback": "true", "X-Error": str(e)[:200]},
        )
